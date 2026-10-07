#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Cập nhật file Excel: Thêm Sheet "MÔI GIỚI & SALES ĐỐI TÁC"
chuyên bán BĐS nghỉ dưỡng, biệt thự biển, đất lớn Đất Đỏ - Hồ Tràm - Bình Châu
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_FILE = "SFR_CRM_KhachHang_Sales_T09_2026.xlsx"

GREEN_DARK   = "183024"
GREEN_LIGHT  = "2C5440"
GOLD         = "C29B53"
GOLD_LIGHT   = "ECDCB9"
GOLD_BG      = "FAF3E6"
WHITE        = "FFFDF9"
BORDER_COLOR = "E2DACB"
BLUE_LIGHT   = "DBEAFE"
GREEN_STATUS = "D1FAE5"

def make_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def make_border(style="thin"):
    s = Side(style=style, color=BORDER_COLOR)
    return Border(left=s, right=s, top=s, bottom=s)

def set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

def col_header(ws, row, col, text, bg=GREEN_DARK, fg="FFFFFF"):
    cell = ws.cell(row=row, column=col, value=text)
    cell.fill  = make_fill(bg)
    cell.font  = Font(name="Calibri", size=9, bold=True, color=fg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = make_border()

def data_cell(ws, row, col, value, align="left", bold=False, bg=None, fg="1F2937", wrap=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font  = Font(name="Calibri", size=9, bold=bold, color=fg)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = make_border()
    if bg:
        cell.fill = make_fill(bg)
    return cell

def title_row(ws, row, col_start, col_end, text, bg=GREEN_DARK, fg="FFFFFF", size=13):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start, value=text)
    cell.fill  = make_fill(bg)
    cell.font  = Font(name="Calibri", size=size, bold=True, color=fg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = make_border()

def sub_header(ws, row, col_start, col_end, text, bg=GOLD, fg=GREEN_DARK, size=10):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start, value=text)
    cell.fill  = make_fill(bg)
    cell.font  = Font(name="Calibri", size=size, bold=True, color=fg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = make_border()

def main():
    wb = openpyxl.load_workbook(EXCEL_FILE)
    
    sheet_name = "🤝 MÔI GIỚI & SALES ĐỐI TÁC"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
        
    ws = wb.create_sheet(sheet_name)
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False

    # Widths
    widths = [5, 6, 22, 16, 22, 22, 22, 20, 28, 25]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 42

    title_row(ws, 2, 2, 10, "🤝 DANH SÁCH SALES & SÀN MÔI GIỚI BĐS NGHỈ DƯỠNG — HỒ TRÀM, BÌNH CHÂU, ĐẤT ĐỎ, BÀ RỊA", size=13)
    sub_header(ws, 3, 2, 10, "Dữ liệu đối tác tiềm năng liên kết bán hàng cho Saigon Farm Resort | Cập nhật 09/2026")

    headers = [
        "#", "STT", "HỌ TÊN / ĐƠN VỊ", "SỐ ĐIỆN THOẠI / ZALO", 
        "VỊ TRÍ / TỔ CHỨC", "ĐỊA BÀN HOẠT ĐỘNG CHÍNH", "PHÂN KHÚC THẾ MẠNH", 
        "ĐỘ PHÙ HỢP VỚI SFR", "TIỀM NĂNG TỆP KHÁCH SẴN CÓ", "GỢI Ý KỊCH BẢN ĐỀ XUẤT HỢP TÁC"
    ]
    for col, h in enumerate(headers, 1):
        col_header(ws, 4, col, h)

    sales_partners = [
        (
            1, "Mỹ Thương Realtor", "0902 390 466",
            "Đông Tây Property (Đông Tây Land)", "Hồ Tràm, Bình Châu, Thủ Đức",
            "Biệt thự biển, Nghỉ dưỡng cao cấp", "⭐⭐⭐⭐⭐ Rất cao",
            "Nắm tệp khách VIP đầu tư BĐS biển triệu USD",
            "Mời làm F1 phân phối, hoa hồng 2.5 - 3.5% + thưởng nóng"
        ),
        (
            2, "Quỳnh Liên", "0936 09 10 11",
            "TPKD - Khải Minh Land", "Hồ Tràm (Maia, Ixora, Sanctuary)",
            "Resort cao cấp, Biệt thự biển 15 - 50 tỷ", "⭐⭐⭐⭐⭐ Rất cao",
            "Đội ngũ sales chuyên đánh dòng sản phẩm Second Home VIP",
            "Ký kết liên kết sàn (Co-broking), tổ chức tour trải nghiệm SFR"
        ),
        (
            3, "Lâm Khánh Tài", "0966 205 205",
            "Beach Villa Vietnam / BĐS Nghỉ dưỡng", "Cung đường ven biển Hồ Tràm - Bình Châu",
            "Villa nghỉ dưỡng, đất view biển, phân tích hạ tầng", "⭐⭐⭐⭐⭐ Rất cao",
            "Kênh YouTube review chuyên nghiệp, tệp nhà đầu tư theo dõi lớn",
            "Tài trợ video review trải nghiệm hồ 100ha + booking độc quyền"
        ),
        (
            4, "Nguyễn Quang Anh", "0909 430 175",
            "Môi giới Đất Lớn Đất Đỏ", "Đất Đỏ, Phước Hội, Láng Dài",
            "Đất mẫu, đất vườn sào, đất kề KCN 1.000ha", "⭐⭐⭐⭐⭐ Cực kỳ sát địa bàn",
            "Nắm toàn bộ khách địa phương và NĐT săn quỹ đất Đất Đỏ",
            "Hợp tác đẩy dòng Điền An & Điền Sản thổ cư có sổ sẵn"
        ),
        (
            5, "Lê Minh Hải", "0938 858 959",
            "Môi giới BĐS Trung tâm Đất Đỏ", "Thị trấn Đất Đỏ, Hồ Lồ Ồ, Long Đất",
            "Đất biệt thự vườn, đất ở ven hồ", "⭐⭐⭐⭐⭐ Cực kỳ sát địa bàn",
            "Khách mua nhà vườn cuối tuần từ TP.HCM đổ về Đất Đỏ",
            "Gửi giỏ hàng Điền An phân khúc 3 - 6 tỷ hoa hồng nhanh"
        ),
        (
            6, "Đạt Nguyễn", "0918 788 966",
            "Vạn Đạt Land (Chuyên Ixora)", "Hồ Tràm Strip, Xuyên Mộc",
            "Biệt thự biển hạng sang, condotel cao cấp", "⭐⭐⭐⭐ Cao",
            "Khách hàng thích nhận dòng tiền cho thuê thụ động",
            "Chào sản phẩm Điền Sản (cam kết vận hành cho thuê turnkey)"
        ),
        (
            7, "Phương Thiện", "0931 555 355",
            "Phương Thiện Land (Chuyên Novaworld)", "Bình Châu, Xuyên Mộc",
            "Shophouse, Biệt thự nghỉ dưỡng Novaworld", "⭐⭐⭐⭐ Cao",
            "Tệp khách Novaworld Hồ Tràm đang tìm sp pháp lý sổ đỏ sở hữu lâu dài",
            "Nhấn mạnh lợi thế: '100% Đất thổ cư lâu dài có sổ đỏ riêng'"
        ),
        (
            8, "Sea Property (Team Hồ Tràm)", "0908 982 299",
            "Công ty CP ĐT BĐS HPR", "Trục Hồ Tràm - Bình Châu",
            "Tổng đại lý phân phối BĐS nghỉ dưỡng biển", "⭐⭐⭐⭐ Cao",
            "Sàn chuyên nghiệp, có phòng marketing và data khách VIP rộng",
            "Ký hợp đồng đối tác chiến lược phân phối Biệt Phủ Điền Trang"
        ),
        (
            9, "Kim Chi Land", "0907 657 687",
            "Sàn BĐS Kim Chi Land", "Xuyên Mộc, Đất Đỏ, Bà Rịa",
            "Đất vườn sinh thái, farmstay, nhà vườn nghỉ dưỡng", "⭐⭐⭐⭐⭐ Rất cao",
            "Tệp khách hàng yêu thích lối sống xanh nông trại (Farm Resort)",
            "Cực kỳ khớp concept văn hóa sinh thái nông trang của SFR"
        ),
        (
            10, "BĐS Vũng Tàu (Mr. Tuấn)", "0978 384 438",
            "Công ty BĐS Vũng Tàu", "Bà Rịa, Châu Đức, Xuyên Mộc",
            "Trang trại, đất vườn lớn, đất ven hồ", "⭐⭐⭐⭐ Cao",
            "Khách tìm đất sào làm nhà vườn ven hồ",
            "Giới thiệu mô hình Điền Trang ven hồ Lồ Ồ 100ha hoàn chỉnh tiện ích"
        ),
        (
            11, "AQLand - Phòng KD Dự án", "0906 462 223",
            "Sàn Giao Dịch AQLand", "Hồ Tràm, Long Hải",
            "Biệt thự biển, phân khúc wellness & retreat", "⭐⭐⭐⭐ Cao",
            "Khách hàng thích dịch vụ chăm sóc sức khỏe, khoáng nóng, thiên nhiên",
            "Pitching concept 5 Trụ cột Lối sống Việt Đương Đại & Wellness ven hồ"
        ),
        (
            12, "CSQ Real Bà Rịa", "0933 666 888",
            "BĐS CSQ Real (Võ Văn Kiệt)", "TP. Bà Rịa, Đất Đỏ, Long Điền",
            "Đất nền thổ cư, nhà phố vườn cao cấp", "⭐⭐⭐⭐ Cao",
            "Khách địa phương giàu có, quan chức và chủ doanh nghiệp tại BR-VT",
            "Tổ chức hội thảo mini giới thiệu Điền Trang cho giới tinh hoa BR-VT"
        )
    ]

    for i, p in enumerate(sales_partners, start=5):
        ws.row_dimensions[i].height = 36
        row = i
        idx = p[0]
        row_bg = WHITE if idx % 2 == 0 else GOLD_BG

        data_cell(ws, row, 1, "", align="center", bg=row_bg)
        data_cell(ws, row, 2, idx, align="center", bold=True, bg=row_bg)
        data_cell(ws, row, 3, p[1], bold=True, bg=row_bg)
        data_cell(ws, row, 4, p[2], align="center", bold=True, bg=GREEN_STATUS, fg="064E3B")
        data_cell(ws, row, 5, p[3], bg=row_bg)
        data_cell(ws, row, 6, p[4], bg=row_bg)
        data_cell(ws, row, 7, p[5], wrap=True, bg=row_bg)
        data_cell(ws, row, 8, p[6], align="center", bold=True, bg=GOLD_LIGHT, fg=GREEN_DARK)
        data_cell(ws, row, 9, p[7], wrap=True, bg=row_bg)
        data_cell(ws, row, 10, p[8], wrap=True, bg=row_bg)

    # Note row
    last_row = len(sales_partners) + 6
    ws.row_dimensions[last_row].height = 55
    ws.merge_cells(start_row=last_row, start_column=2, end_row=last_row, end_column=10)
    note_cell = ws.cell(row=last_row, column=2,
        value="📌 CHIẾN LƯỢC TIẾP CẬN MÔI GIỚI: 1) Liên hệ trực tiếp qua Zalo/Điện thoại gửi Bảng hoa hồng & Cơ chế chia sẻ hấp dẫn (2.5% - 3.5%). "
              "2) Nhấn mạnh các điểm 'bán được ngay': Sổ đỏ thổ cư riêng từng nền, mặt hồ nước ngọt 100ha độc bản, giá chỉ 1/3 đất biển Hồ Tràm. "
              "3) Mời tham quan thực tế trải nghiệm tại Saigon Farm Resort và tặng voucher trải nghiệm ẩm thực Farm-to-Table.")
    note_cell.fill = make_fill(GREEN_DARK)
    note_cell.font = Font(name="Calibri", size=9.5, bold=True, color="FFFDF9")
    note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    note_cell.border = make_border()

    wb.save(EXCEL_FILE)
    print(f"✅ Đã thêm sheet '{sheet_name}' vào file {EXCEL_FILE} thành công!")

if __name__ == "__main__":
    main()
